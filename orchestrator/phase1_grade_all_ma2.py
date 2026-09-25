"""
phase1_grade_all_ma2.py — MA2 Grading orchestrator

Purpose: Grades the formula-based parts of every student's MA2 workbook.
         Handles Income and Projection, Credit Cards, Student Loans,
         and Annual Budget tabs.

Author: Clayton Ragsdale
"""

import os
from openpyxl import load_workbook
from typing import Dict, Any, Optional

from utilities.logger import get_logger

# MA2 Graders
from graders.ma2_income_projection.grade_income_projection_tab import grade_income_projection_tab
from graders.ma2_credit_cards.grade_credit_card_tab import grade_credit_cards_tab
from graders.ma2_student_loans.grade_student_loans_tab import grade_student_loans_tab
from graders.ma2_annual_budget.grade_annual_budget_tab import grade_annual_budget_tab

# MA2 Writers
from writers.write_ma2_income_projection_results import write_income_projection_results
from writers.write_ma2_credit_card_results import write_credit_card_results_to_grading_sheet
from writers.write_ma2_student_loans_results import write_student_loans_results_to_grading_sheet
from writers.write_ma2_annual_budget_results import write_annual_budget_results_to_grading_sheet

# ---- Error Handling ----
from utilities.errors import GradingResult, GradingError, ErrorCategory, classify_error
# ------------------------


def _validate_ma2_sheets(workbook) -> tuple:
    """
    Validate that required MA2 sheets exist (case-insensitive).

    Returns:
        Tuple of (is_valid, sheet_map, missing_sheets)
        sheet_map maps canonical names to actual sheet names
    """
    required_sheets = ["Income and Projection", "Credit Cards", "Student Loans", "Annual Budget"]
    sheet_names_lower = {s.lower(): s for s in workbook.sheetnames}

    sheet_map = {}
    missing_sheets = []

    for req in required_sheets:
        if req.lower() in sheet_names_lower:
            sheet_map[req] = sheet_names_lower[req.lower()]
        else:
            missing_sheets.append(req)

    is_valid = len(missing_sheets) == 0
    return is_valid, sheet_map, missing_sheets


def phase1_grade_all_students_ma2(
    submissions_path: str,
    graded_output_path: str,
    pipeline_state: Optional[Dict[str, Any]] = None
) -> None:
    """
    Grades the formula-based parts of every student's MA2 workbook.

    Args:
        submissions_path: Path to folder containing student submission files
        graded_output_path: Path to folder containing grading sheet templates
        pipeline_state: Optional dict for cancellation checking
    """
    logger = get_logger()
    logger.info("")
    logger.info("=" * 60)
    logger.info("PHASE 1 - Grading all MA2 students...")
    logger.info("=" * 60)

    student_files = [f for f in os.listdir(submissions_path) if f.endswith(".xlsx")]
    # Single-submission mode: grade only one file (used by the "late student" flow).
    _only = pipeline_state.get("only_filename") if pipeline_state else None
    if _only:
        student_files = [f for f in student_files if f == _only]
    total_students = len(student_files)
    logger.info(f"Found {total_students} student submissions to grade")

    # Initialize result tracker
    result = GradingResult()
    skipped_count = 0

    # Store result in pipeline_state for frontend access
    if pipeline_state is not None:
        pipeline_state["grading_result"] = result

    for idx, filename in enumerate(student_files, 1):
        # Check for cancellation request
        if pipeline_state and pipeline_state.get("cancel_requested"):
            logger.warning(f"Pipeline cancelled by user after grading {idx - 1} students")
            break

        # Extract student name from filename
        student_name = filename.replace("_MA2.xlsx", "").replace(".xlsx", "")
        submission_file = os.path.join(submissions_path, filename)
        grading_file = os.path.join(graded_output_path, f"{student_name}_MA2_Grade.xlsx")

        logger.info(f"[{idx}/{total_students}] Processing: {student_name}")

        # Emit live progress to any subscriber (e.g. the desktop app)
        if pipeline_state:
            _cb = pipeline_state.get("progress_callback")
            if _cb:
                try:
                    _cb(idx, total_students, student_name)
                except Exception:
                    pass

        student_wb = None
        grading_wb = None

        try:
            logger.debug(f"  Loading submission: {submission_file}")
            student_wb = load_workbook(submission_file, data_only=False)
            student_wb_data = load_workbook(submission_file, data_only=True)  # calculated values

            logger.debug(f"  Loading grading sheet: {grading_file}")
            grading_wb = load_workbook(grading_file)

            ws_grading = grading_wb["Grading Sheet"]

            # Validate required sheets
            is_valid, sheet_map, missing_sheets = _validate_ma2_sheets(student_wb)

            if missing_sheets:
                logger.warning(f"  Missing sheets for {student_name}: {missing_sheets}")

            # -----------------------------
            # INCOME AND PROJECTION TAB
            # (grader takes worksheet object)
            # -----------------------------
            if sheet_map.get("Income and Projection"):
                try:
                    logger.debug(f"  Grading Income and Projection tab...")
                    ws_income = student_wb[sheet_map["Income and Projection"]]
                    ws_income_data = student_wb_data[sheet_map["Income and Projection"]]
                    income_results = grade_income_projection_tab(ws_income, ws_income_data)
                    write_income_projection_results(ws_grading, income_results)
                    logger.debug(f"  Income and Projection tab complete")
                except Exception as e:
                    logger.warning(f"  Income and Projection error for {student_name}: {e}")
            else:
                logger.info(f"  Skipping Income and Projection (sheet missing)")
                skipped_count += 1

            # -----------------------------
            # CREDIT CARDS TAB
            # (grader takes file path)
            # (writer takes workbook object)
            # -----------------------------
            if sheet_map.get("Credit Cards"):
                try:
                    logger.debug(f"  Grading Credit Cards tab...")
                    credit_results = grade_credit_cards_tab(submission_file)
                    write_credit_card_results_to_grading_sheet(grading_wb, credit_results)
                    logger.debug(f"  Credit Cards tab complete")
                except Exception as e:
                    logger.warning(f"  Credit Cards error for {student_name}: {e}")
            else:
                logger.info(f"  Skipping Credit Cards (sheet missing)")
                skipped_count += 1

            # -----------------------------
            # STUDENT LOANS TAB
            # (grader takes file path)
            # (writer takes workbook object)
            # -----------------------------
            if sheet_map.get("Student Loans"):
                try:
                    logger.debug(f"  Grading Student Loans tab...")
                    loans_results = grade_student_loans_tab(submission_file)
                    write_student_loans_results_to_grading_sheet(grading_wb, loans_results)
                    logger.debug(f"  Student Loans tab complete")
                except Exception as e:
                    logger.warning(f"  Student Loans error for {student_name}: {e}")
            else:
                logger.info(f"  Skipping Student Loans (sheet missing)")
                skipped_count += 1

            # -----------------------------
            # ANNUAL BUDGET TAB
            # (grader takes file path)
            # (writer takes workbook object)
            # -----------------------------
            if sheet_map.get("Annual Budget"):
                try:
                    logger.debug(f"  Grading Annual Budget tab...")
                    budget_results = grade_annual_budget_tab(submission_file)
                    write_annual_budget_results_to_grading_sheet(grading_wb, budget_results)
                    logger.debug(f"  Annual Budget tab complete")
                except Exception as e:
                    logger.warning(f"  Annual Budget error for {student_name}: {e}")
            else:
                logger.info(f"  Skipping Annual Budget (sheet missing)")
                skipped_count += 1

            grading_wb.save(grading_file)
            logger.info(f"  ✓ Graded: {student_name}")
            result.add_success()

        except Exception as e:
            grading_error = classify_error(e, student_name=student_name)
            result.add_issue(grading_error)
            logger.error(f"  ✗ {grading_error.problem}")

        finally:
            # Ensure workbooks are always closed
            if student_wb is not None:
                student_wb.close()
            if grading_wb is not None:
                grading_wb.close()

    # Summary
    logger.info("")
    logger.info("-" * 40)
    issue_count = len(result.issues)
    logger.info(f"Phase 1 Summary: {result.success_count} graded, {issue_count} issues, {skipped_count} sheets skipped")

    # Log any student issues
    if result.issues:
        logger.info("")
        logger.info("Student Issues:")
        for issue in result.issues:
            name = issue.student_name or "Unknown"
            logger.info(f"  • {name}: {issue.problem}")
            logger.info(f"    → {issue.action}")

    logger.info("-" * 40)
