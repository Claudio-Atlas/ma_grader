"""
check_predictions_formatting.py — Validate currency formatting on predicted values

Purpose: Checks that the predicted salary values in E19:E35 are formatted as
         currency with 0 decimal places. Since these are salary predictions,
         they should display as whole dollar amounts (e.g., "$52,000" not "$52,000.00").

Author: Clayton Ragsdale
Dependencies: None (uses only openpyxl worksheet objects)

Expected Format:
    Currency format with 0 decimal places, such as:
    - "$#,##0" or "$#,##0_"
    - "$0" or similar variants
"""

from typing import Tuple, List, Dict, Any
from openpyxl.worksheet.worksheet import Worksheet


def _is_currency_zero_decimal(fmt: str) -> bool:
    """
    Check if a number format is currency with 0 decimal places.
    
    This function validates Excel number format strings to determine if they
    display values as currency without decimal places.
    
    Validation Rules:
        1. Must contain "$" (currency indicator)
        2. Must NOT have decimal places (no ".0" or ".#" patterns)
        3. Must have number placeholders (0 or #)
    
    Args:
        fmt: The Excel number format string to validate
    
    Returns:
        bool: True if format is currency with 0 decimals, False otherwise
    
    Example:
        >>> _is_currency_zero_decimal("$#,##0")
        True
        >>> _is_currency_zero_decimal("$#,##0.00")
        False
        >>> _is_currency_zero_decimal("#,##0")  # No currency symbol
        False
    """
    if not fmt or not isinstance(fmt, str):
        return False
    
    # Rule 1: Must contain $ somewhere (currency indicator)
    if "$" not in fmt:
        return False
    
    # Rule 2: Should NOT have decimal places (no .0, .00, .#, etc.)
    if ".0" in fmt or ".#" in fmt:
        return False
    
    # Rule 3: Must have number placeholders (0 or #)
    if "0" not in fmt and "#" not in fmt:
        return False
    
    return True


def check_currency_formatting(ws: Worksheet) -> Tuple[float, List[Tuple[str, Dict[str, Any]]]]:
    """
    Check formatting of predicted salary values in E19:E35.
    
    Expected Formatting:
        All cells should be formatted as currency with 0 decimal places.
        This is appropriate for salary predictions which should display
        as whole dollar amounts.
    
    Grading Logic:
        - 1 point total for all 17 cells (E19:E35)
        - Points are proportional: (correct_count / 17)
        - Full credit requires ALL cells to be formatted correctly
    
    Args:
        ws: The openpyxl Worksheet object for the Income Analysis tab
    
    Returns:
        Tuple of:
            - score: float (0.0 to 1.0)
            - feedback: List of (code, params) tuples describing results
    
    Example:
        >>> score, feedback = check_currency_formatting(worksheet)
        >>> print(f"Formatting score: {score}/1.0")
    """
    total_rows = 17  # E19 through E35
    correct_count = 0
    sample_incorrect_fmt = None  # Store first incorrect format for feedback

    # Check each prediction cell
    for row in range(19, 36):  # Rows 19-35 inclusive
        cell = ws[f"E{row}"]
        fmt = cell.number_format
        
        if _is_currency_zero_decimal(fmt):
            correct_count += 1
        elif sample_incorrect_fmt is None:
            # Store the first incorrect format we find (for feedback)
            sample_incorrect_fmt = fmt

    # ============================================================
    # Build Feedback
    # ============================================================
    
    # All correctly formatted
    if correct_count == total_rows:
        return 1.0, [("IA_PRED_FORMAT_ALL_CORRECT", {"range": "E19:E35"})]

    # None correctly formatted
    if correct_count == 0:
        return 0.0, [("IA_PRED_FORMAT_NONE_CORRECT", {
            "range": "E19:E35",
            "found": sample_incorrect_fmt
        })]

    # Partial credit - proportional scoring
    score = round(correct_count / total_rows, 2)
    return score, [("IA_PRED_FORMAT_PARTIAL", {
        "correct": correct_count,
        "total": total_rows,
        "range": "E19:E35"
    })]
