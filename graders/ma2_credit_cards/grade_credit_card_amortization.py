"""
grade_credit_card_amortization.py — Credit Cards amortization table (24 pts).

The table (starting at row 25) has one row per month with five columns:

    A: Month number        (1, 2, 3, ...)
    B: Current Balance      B25 = C14 (start balance);  B{r} = D{r-1} + E{r-1}
    C: Payment              = MIN(B{r}, MAX(C18, C17*B{r}))
    D: Balance After Pay    = B{r} - C{r}
    E: Interest             = D{r} * (C13/C15)

Scoring (per Richard's suggestion): the four computed columns B, C, D, E are each
worth 6 points (24 total), scored by how many rows follow the correct formula
pattern. This is robust to how far a student drags the table and to where it pays
off — the formula pattern is identical on every row — and it gives clear,
per-column feedback instead of a single opaque number.
"""

from openpyxl import load_workbook

START_ROW = 25
MAX_SCAN = 800  # safety cap
COL_LABELS = {
    "B": "Current Balance",
    "C": "Payment",
    "D": "Balance After Payment",
    "E": "Interest",
}


def _norm(formula) -> str:
    if not isinstance(formula, str):
        return ""
    return formula.strip().lower().replace(" ", "").replace("$", "")


def _find_table_rows(ws_f):
    """Return the list of populated table rows: START_ROW down to the last row
    whose Payment (C) cell holds a formula. Stops after a few blank rows."""
    last = START_ROW
    blanks = 0
    r = START_ROW
    while r < START_ROW + MAX_SCAN:
        c = ws_f[f"C{r}"].value
        if isinstance(c, str) and c.startswith("="):
            last = r
            blanks = 0
        else:
            blanks += 1
            if blanks >= 3:
                break
        r += 1
    return list(range(START_ROW, last + 1)), last


def _row_ok(col, r, B, C, D, E):
    """Is the formula in `col` at row r the correct pattern?"""
    if col == "B":
        if r == START_ROW:
            # Start balance: references C14 (often wrapped in an IF(... , C14)).
            return "c14" in B
        return (f"d{r-1}" in B) and (f"e{r-1}" in B) and ("+" in B)
    if col == "C":
        return ("min(" in C) and ("max(" in C) and (f"b{r}" in C) and ("c17" in C) and ("c18" in C)
    if col == "D":
        return (f"b{r}" in D) and (f"c{r}" in D) and ("-" in D)
    if col == "E":
        return (
            (f"d{r}" in E) and ("c13" in E)
            and (("c15" in E) or ("/12" in E)) and ("/" in E) and ("*" in E)
        )
    return False


def _detect_payoff(ws_v, rows):
    """Payoff row = first row where the Balance After Payment (D) reaches ~0 with
    a positive starting balance. Falls back to the last row."""
    for r in rows:
        b = ws_v[f"B{r}"].value
        d = ws_v[f"D{r}"].value
        if isinstance(b, (int, float)) and isinstance(d, (int, float)):
            if b > 0.005 and abs(d) <= 0.05:
                return r
    return rows[-1] if rows else START_ROW


def grade_credit_card_amortization(student_file_path):
    """
    Grade the amortization table. Returns:
        {
            "total_points": float (0-24),
            "comment": str (per-column breakdown),
            "column_points": {"B":.., "C":.., "D":.., "E":..},
            "payoff_row": {"payoff_month": int, "excel_row": int, "last_row": int},
        }
    """
    wb_f = load_workbook(student_file_path, data_only=False)
    wb_v = load_workbook(student_file_path, data_only=True)
    ws_f = wb_f["Credit Cards"]
    ws_v = wb_v["Credit Cards"]

    rows, last = _find_table_rows(ws_f)
    n = len(rows)

    correct = {c: 0 for c in "BCDE"}
    bad = {c: [] for c in "BCDE"}

    for r in rows:
        B = _norm(ws_f[f"B{r}"].value)
        C = _norm(ws_f[f"C{r}"].value)
        D = _norm(ws_f[f"D{r}"].value)
        E = _norm(ws_f[f"E{r}"].value)
        vals = {"B": B, "C": C, "D": D, "E": E}
        for col in "BCDE":
            if _row_ok(col, r, B, C, D, E):
                correct[col] += 1
            else:
                bad[col].append(f"{col}{r}")

    column_points = {}
    comments = []
    total = 0.0
    for col in "BCDE":
        pts = round(6.0 * correct[col] / n, 2) if n else 0.0
        column_points[col] = pts
        total += pts
        if correct[col] == n:
            comments.append(f"{COL_LABELS[col]} column ({col}): all {n} rows correct — {pts}/6.")
        else:
            sample = ", ".join(bad[col][:4]) + (" …" if len(bad[col]) > 4 else "")
            comments.append(
                f"{COL_LABELS[col]} column ({col}): {correct[col]}/{n} rows follow the "
                f"correct formula — {pts}/6. First issues: {sample}."
            )

    total = round(total, 2)

    payoff_row = _detect_payoff(ws_v, rows)
    a_val = ws_v[f"A{payoff_row}"].value
    payoff_month = int(a_val) if isinstance(a_val, (int, float)) else (payoff_row - START_ROW + 1)

    wb_f.close()
    wb_v.close()

    return {
        "total_points": total,
        "comment": " ".join(comments),
        "column_points": column_points,
        "payoff_row": {"payoff_month": payoff_month, "excel_row": payoff_row, "last_row": last},
    }
