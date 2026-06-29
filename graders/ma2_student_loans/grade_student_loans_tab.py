from openpyxl import load_workbook

from graders.ma2_student_loans.grade_mortgage_rate_lookup import grade_mortgage_rate_lookup
from graders.ma2_student_loans.grade_payment_parameters import grade_payment_parameters
from graders.ma2_student_loans.grade_subsidized_payment_section import grade_subsidized_payment_section
from graders.ma2_student_loans.grade_unsubsidized_setup import grade_unsubsidized_setup
from graders.ma2_student_loans.grade_unsubsidized_payment_section import grade_unsubsidized_payment_section
from graders.ma2_student_loans.grade_student_loans_formatting import grade_student_loans_formatting


def grade_student_loans_tab(student_file_path):
    """
    Master wrapper for the entire Student Loans tab.
    Returns all section dictionaries with NO total score.
    """

    try:
        # Load both versions of workbook
        wb_f = load_workbook(student_file_path, data_only=False)   # formulas
        wb_v = load_workbook(student_file_path, data_only=True)    # values

        ws_f = wb_f["Student Loans"]
        ws_v = wb_v["Student Loans"]

        # -----------------------------
        # Run all sub-graders
        # -----------------------------
        mortgage_lookup = grade_mortgage_rate_lookup(ws_v)

        payment_params = grade_payment_parameters(student_file_path)

        subsidized_payments = grade_subsidized_payment_section(student_file_path)

        unsub_setup = grade_unsubsidized_setup(ws_f)

        unsub_payments = grade_unsubsidized_payment_section(student_file_path)

        formatting = grade_student_loans_formatting(ws_f)

        # -----------------------------
        # Assemble full results
        # -----------------------------
        results = {
            "mortgage_rate_lookup": mortgage_lookup,
            "payment_parameters": payment_params,
            "subsidized_payments": subsidized_payments,
            "unsubsidized_setup": unsub_setup,
            "unsubsidized_payments": unsub_payments,
            "formatting": formatting,
        }

        wb_f.close()
        wb_v.close()

        return results

    except Exception as e:
        return {"error": f"Error grading Student Loans tab: {e}"}
