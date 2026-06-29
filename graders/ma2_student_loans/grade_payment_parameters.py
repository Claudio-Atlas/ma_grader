from openpyxl import load_workbook

def grade_payment_parameters(student_file_path):
    """
    Grades the simple input values in B27 and B28 on the 'Student Loans' tab.
    Each cell is worth 1 point.

    Expected defaults:
        B27 = 12   (payments per year)
        B28 = 10   (loan years)

    Returns a dictionary with points, comments, and total.
    """
    # --- Editable expected values ---
    expected_values = {
        "B27": 12,
        "B28": 10,
    }

    wb = load_workbook(student_file_path, data_only=True)
    ws = wb["Student Loans"]

    results = {}
    total_points = 0

    for cell, expected in expected_values.items():
        student_val = ws[cell].value
        entry = {"points": 0, "comment": ""}

        if student_val == expected:
            entry["points"] = 1
            entry["comment"] = f"{cell} value correct ({student_val})."
        else:
            entry["comment"] = (
                f"{cell} value incorrect. Expected {expected}, found {student_val}."
            )

        total_points += entry["points"]
        results[cell] = entry

    wb.close()

    results["total_points"] = total_points
    results["comment"] = " | ".join([r["comment"] for c, r in results.items() if c in expected_values])

    return results
