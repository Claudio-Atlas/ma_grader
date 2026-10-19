"""
xl.py — load student workbooks with Excel's "future function" prefixes removed.

Excel stores functions added after 2007 with a hidden prefix in the file, e.g.
STDEV.S(...) is saved as _xlfn.STDEV.S(...), and dynamic-array helpers as
_xlfn._xlws.SORT(...). Graders compare formula text, so the prefix made correct
formulas look wrong (=AVERAGE(D14:D63)-STDEV.S(D14:D63) was rejected).

Only use this for workbooks that are read, never saved: Excel needs the
prefixes in files it opens.
"""

import re

from openpyxl import load_workbook as _openpyxl_load

_PREFIX_RE = re.compile(r"_xl(?:fn|ws|udf|pm)\.", re.I)


def strip_future_prefixes(wb):
    """Remove _xlfn./_xlws./_xludf. prefixes from every formula in `wb`."""
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if isinstance(v, str):
                    if v.startswith("=") and "_xl" in v.lower():
                        cell.value = _PREFIX_RE.sub("", v)
                elif type(v).__name__ == "ArrayFormula":
                    text = getattr(v, "text", None)
                    if isinstance(text, str) and "_xl" in text.lower():
                        v.text = _PREFIX_RE.sub("", text)
    return wb


def load_workbook(*args, **kwargs):
    """openpyxl.load_workbook for STUDENT files: formulas come back without the
    hidden future-function prefixes. Cached values (data_only=True) are untouched."""
    wb = _openpyxl_load(*args, **kwargs)
    if not kwargs.get("data_only"):
        strip_future_prefixes(wb)
    return wb
