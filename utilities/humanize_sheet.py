"""
humanize_sheet.py — final pass that rewrites a graded sheet's Comments column
in the student-facing voice (see student_feedback.py) and adds an overall note.

It runs after every tab writer, reads each rubric row's points (F), possible
points (E) and the grader's comment (G), and rewrites G. Graders that report
problems in the new voice are used as-is; older grader comments are filtered
down to just the problems.
"""

import math

from openpyxl.styles import Alignment, Font

from utilities import student_feedback as sf

FORMATTING_PRAISE = "Your formatting on this tab is spot on. Every cell is set up the way the instructions ask."

# Per-row praise for full credit, keyed by grading-sheet row.
MA2_PRAISE = {
    3: "Thanks for entering your full name.",
    4: "Your CPI table is complete, and every value matches the reference data. Nice work!",
    5: "Your SLOPE and INTERCEPT formulas are exactly right.",
    6: "Your projection year, projected CPI and inflation rate are all correct. Great job.",
    7: "Your current income, 5-year projection and monthly income formulas are all correct.",
    8: FORMATTING_PRAISE,
    10: "You pulled the right interest rate from the Mortgage Rates tab.",
    11: "Your unsubsidized loan setup (years in school, added interest and new principal) is all correct.",
    12: "Payments per year and number of years are both correct.",
    13: "Your subsidized loan payment, total paid and interest paid are all correct. Nice work!",
    14: "Your unsubsidized loan payment, total paid and interest paid are all correct.",
    15: FORMATTING_PRAISE,
    17: "All five credit card inputs match the values in the description.",
    18: "Your amortization table is correct all the way to payoff. Great job.",
    19: "Years to pay off, total paid and total interest are all calculated correctly.",
    20: FORMATTING_PRAISE,
    22: "Your Housing through Other rows are complete, with a formula for every Total Annual Cost.",
    23: "Your subsidized loan row is correct and pulls the payment straight from the Student Loans tab.",
    24: "Your credit card row is correct and pulls the payment straight from the Credit Cards tab.",
    25: "Your Total Annual Budget adds everything up with SUM, just as asked.",
    26: "Your 5-year projected total and remaining total are both correct.",
    27: "Your pie chart has everything: the right data, a category legend, a clear title and percentage labels. Nice work!",
    28: FORMATTING_PRAISE,
}

# Short names for the biggest pieces of work, used to praise strengths in the
# overall note (only rows earning full credit are mentioned).
MA2_STRENGTHS = {
    18: "credit card amortization table",
    22: "budget table",
    13: "subsidized loan calculations",
    14: "unsubsidized loan calculations",
    4: "CPI table",
    6: "CPI projection",
    7: "income projection",
    27: "pie chart",
    19: "payoff summary",
}

MA2_SECTIONS = [
    ("Income and Projection", range(3, 9)),
    ("Student Loans", range(10, 16)),
    ("Credit Cards", range(17, 21)),
    ("Annual Budget", range(22, 29)),
]

MA1_PRAISE = {
    3: "Thanks for entering your full name.",
    4: "Your slope and y-intercept are calculated correctly.",
    5: "Your Average Weekly Income formulas are all correct and use cell references.",
    6: "Your XY scatterplot is set up exactly as asked.",
    7: "Your trendline is in place and extended as asked.",
    9: "Your final quantities all have the right units.",
    10: "You picked the right conversion ratios, with the right units.",
    11: "Your ratio formulas are correct and use cell references.",
    12: "Your final quantity formulas are all correct.",
    13: "Your Fahrenheit-to-Celsius formulas are correct.",
    15: "Your name letters are entered correctly.",
    16: "Your countries are good choices from the list.",
    17: "Your lookup dates are all filled in and consistent with each other.",
    18: "Your currency codes are all correct.",
    19: "Your exchange rates are all entered to enough significant digits.",
    20: "Your foreign-currency budget formulas are correct.",
    21: "Your conversions back to dollars are correct.",
    22: "Your date, rate, budget and dollar cells are all formatted correctly.",
}
MA1_STRENGTHS = {
    11: "ratio formulas", 12: "final quantity formulas", 20: "foreign-currency budget",
    21: "conversions back to dollars", 4: "best-fit line", 5: "weekly income formulas",
    10: "conversion ratios", 6: "scatterplot",
}
MA1_SECTIONS = [
    ("Income Analysis", range(3, 8)),
    ("Unit Conversions", range(9, 14)),
    ("Currency Conversion", range(15, 23)),
]

MA3_PRAISE = {
    3: "Thanks for entering your full name.",
    4: "Your differences (After minus Before) are all calculated correctly.",
    5: "Your mean, median, standard deviation and range are all correct.",
    6: "Your percentile formulas are correct.",
    7: "Your empirical-rule bounds are correct.",
    # 8 = Written Analysis: graded by hand, left exactly as written.
    9: FORMATTING_PRAISE,
    11: "Your bin minimum, maximum and width are all correct.",
    12: "Your lower and upper limits are correct all the way down the table.",
    13: "Your Title of Bin, Frequency and Relative Frequency formulas are all correct.",
    14: "Your histogram is complete: the right data, touching bars and clear titles. Nice work!",
    15: FORMATTING_PRAISE,
}
MA3_STRENGTHS = {
    5: "summary statistics", 13: "frequency distribution", 12: "bin limits",
    4: "difference calculations", 7: "empirical-rule bounds", 14: "histogram", 6: "percentiles",
}
MA3_SECTIONS = [
    ("Analysis", range(3, 10)),
    ("Visualization", range(11, 16)),
]

LAYOUTS = {
    # label cell, note cell, merged range for the note
    "MA1": {"praise": MA1_PRAISE, "sections": MA1_SECTIONS, "strengths": MA1_STRENGTHS,
            "label": "C28", "note": "C29", "merge": "C29:G33"},
    "MA2": {"praise": MA2_PRAISE, "sections": MA2_SECTIONS, "strengths": MA2_STRENGTHS,
            "label": "G31", "note": "G32", "merge": "G32:H37"},
    "MA3": {"praise": MA3_PRAISE, "sections": MA3_SECTIONS, "strengths": MA3_STRENGTHS,
            "label": "C21", "note": "C22", "merge": "C22:G26"},
}


LINE_PT = 15.5  # points per wrapped line at the template's 11pt font


def _lines(text, width_chars):
    return sum(max(1, math.ceil(len(part) / max(10, width_chars))) for part in str(text).split("\n"))


def _fit_row(ws, row, text, col="G"):
    """Grow a fixed-height row so a wrapped comment isn't cut off (never shrinks)."""
    width = (ws.column_dimensions[col].width or 50) * 1.05
    need = _lines(text, width) * LINE_PT + 4
    current = ws.row_dimensions[row].height or 15
    if need > current:
        ws.row_dimensions[row].height = need


def humanize_grading_sheet(ws, assignment, student_name="", row_issues=None):
    """Rewrite column G for every rubric row and add the overall note.

    row_issues: optional {row: [student-facing issue, ...]} from graders that
    report problems directly; those replace the legacy comment for that row.
    Returns the overall note text (or None if the layout is unknown).
    """
    layout = LAYOUTS.get(assignment.upper())
    if not layout:
        return None
    row_issues = row_issues or {}

    num = sf._num

    def hand_graded(row):
        g = ws[f"G{row}"].value
        return isinstance(g, str) and g.lstrip().startswith("[MANUAL GRADING")

    manual_rows = {r for _, rows in layout["sections"] for r in rows if hand_graded(r)}
    comments = []
    for row, praise in layout["praise"].items():
        pts, poss = num(ws[f"F{row}"].value), num(ws[f"E{row}"].value)
        if poss is None:
            continue
        if row in manual_rows:
            continue  # hand-graded row: leave the instructor's material alone
        legacy = ws[f"G{row}"].value
        # Notes meant for the instructor stay visible, below the student text.
        legacy, notes = sf.split_instructor_notes(legacy)
        text = sf.compose(praise, pts, poss, row_issues.get(row), fallback=legacy)
        if notes:
            cleaned = [n.replace("[MANUAL REVIEW]", "").strip() for n in notes]
            cleaned = [n if n.lower().startswith("instructor check") else "Instructor check: " + n for n in cleaned]
            text += "\n\n" + " ".join(cleaned)
        ws[f"G{row}"] = text
        ws[f"G{row}"].alignment = Alignment(wrap_text=True, vertical="top")
        _fit_row(ws, row, text)
        comments.append(text)

    sections = []
    for name, rows in layout["sections"]:
        # Hand-graded rows aren't scored yet, so they'd look like lost points.
        rows = [r for r in rows if r not in manual_rows]
        pts = sum(v for r in rows if (v := num(ws[f"F{r}"].value)) is not None)
        poss = sum(v for r in rows if (v := num(ws[f"E{r}"].value)) is not None)
        if poss:
            sections.append({"name": name, "points": pts, "possible": poss})

    # Biggest fully-correct rows first.
    def full(row):
        e, f = num(ws[f"E{row}"].value), num(ws[f"F{row}"].value)
        return e is not None and f is not None and f >= e - 1e-9

    strengths = [
        phrase for row, phrase in sorted(layout.get("strengths", {}).items(),
                                         key=lambda kv: -(num(ws[f"E{kv[0]}"].value) or 0))
        if full(row)
    ]

    first = sf.first_name(student_name)
    note = None
    if sf.ai_notes_enabled():
        note = sf.ai_overall_note(first, sections, comments)
    if not note:
        note = sf.overall_note(first, sections, strengths)

    ws[layout["label"]] = "Overall feedback"
    ws[layout["label"]].font = Font(bold=True)
    try:
        ws.merge_cells(layout["merge"])
    except ValueError:
        pass  # already merged (re-grade)
    ws[layout["note"]] = note
    ws[layout["note"]].alignment = Alignment(wrap_text=True, vertical="top")

    # Make the merged note box tall enough: extra height goes on its last row.
    from openpyxl.utils import range_boundaries
    c1, r1, c2, r2 = range_boundaries(layout["merge"])
    from openpyxl.utils import get_column_letter
    width = sum((ws.column_dimensions[get_column_letter(c)].width or 10) for c in range(c1, c2 + 1)) * 1.05
    have = sum((ws.row_dimensions[r].height or 15) for r in range(r1, r2 + 1))
    need = _lines(note, width) * LINE_PT + 6
    if need > have:
        ws.row_dimensions[r2].height = (ws.row_dimensions[r2].height or 15) + (need - have)
    return note
