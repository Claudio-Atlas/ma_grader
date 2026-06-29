import os
from openpyxl import load_workbook

def write_student_loans_results_to_grading_sheet(
    wb_grading,
    results: dict
):
    """
    Writes the Student Loans tab grading results (Major Assignment 2)
    into the existing student's grading sheet workbook.

    PARAMETERS
    ----------
    wb_grading : openpyxl.Workbook
        Already-opened grading workbook (the student's personalized sheet).
    results : dict
        Output from grade_student_loans_tab(student_file_path)

    Writes points into column F and comments into column G for rows 10–15.

    Row Mapping (confirmed by user):
        F10/G10 → Mortgage Rate Lookup (1 pt)
        F11/G11 → Unsubsidized Setup (5 pts)
        F12/G12 → Payment Parameters (2 pts)
        F13/G13 → Subsidized Payments (8 pts)
        F14/G14 → Unsubsidized Payments (8 pts)
        F15/G15 → Formatting (3 pts)

    Does NOT overwrite subtotal row (Row 16).
    """

    ws = wb_grading["Grading Sheet"]

    # -------------------------------------------------------
    # Extract all results from wrapper output
    # -------------------------------------------------------
    mortgage = results.get("mortgage_rate_lookup", {})
    unsub_setup = results.get("unsubsidized_setup", {})
    payment_params = results.get("payment_parameters", {})
    subsidized = results.get("subsidized_payments", {})
    unsub_payments = results.get("unsubsidized_payments", {})
    formatting = results.get("formatting", {})

    # -------------------------------------------------------
    # Safe getters for points + comments
    # -------------------------------------------------------
    def pts(section):
        return section.get("total_points", 0)

    def cmt(section):
        return section.get("comment", "")

    # -------------------------------------------------------
    # Write into Grading Sheet
    # -------------------------------------------------------
    ws["F10"] = pts(mortgage)
    # The mortgage grader nests its message under "B14"; fall back to that.
    ws["G10"] = cmt(mortgage) or mortgage.get("B14", {}).get("comment", "")

    ws["F11"] = pts(unsub_setup)
    ws["G11"] = cmt(unsub_setup)

    ws["F12"] = pts(payment_params)
    ws["G12"] = cmt(payment_params)

    ws["F13"] = pts(subsidized)
    ws["G13"] = cmt(subsidized)

    ws["F14"] = pts(unsub_payments)
    ws["G14"] = cmt(unsub_payments)

    ws["F15"] = pts(formatting)
    ws["G15"] = cmt(formatting)

    # (Do not touch F16/G16 — subtotal row handled by Excel)

    return wb_grading
