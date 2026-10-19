#!/usr/bin/env python3
"""
Regression checks for grading issues instructors reported (v1.4 issue lists).

Each case is a formula that was graded wrong; this asserts the corrected
behavior. Run from the engine root:

    python3 tests/reported_issues.py
"""

import os
import sys

_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if _ROOT not in sys.path:
    sys.path.insert(0, _ROOT)

from openpyxl import Workbook

from utilities.xl import strip_future_prefixes
from graders.currency_conversion.row20_budget_conversion_v2 import grade_row20_budget_conversion_v2
from graders.ma3_analysis.check_empirical_rule import _bound_credit, check_empirical_rule
from graders.ma3_analysis.check_written_analysis import check_written_analysis
from graders.ma3_visualization.check_bin_table import _check_min_formula
from graders.ma3_visualization.check_freq_dist import (
    _check_frequency_formula, _check_relative_freq_formula, check_freq_dist_values,
)

failures = []


def check(label, got, want):
    ok = got == want
    print(f"  {'ok  ' if ok else 'FAIL'} {label}: got {got!r}, want {want!r}")
    if not ok:
        failures.append(label)


print("MA1 · Currency Conversion C20:F20")
for label, tmpl, want in [
    ("text string =\"MXN \"&(C19*B4) gets no credit", '="MXN "&({r}*B4)', 0.0),
    ("=B4*C19 full credit", "=B4*{r}", 8.0),
    ("=ROUND(B4*C19,2) full credit", "=ROUND(B4*{r},2)", 8.0),
]:
    ws = Workbook().active
    ws["B4"] = 500
    for c in "CDEF":
        ws[f"{c}19"] = 20.1
        ws[f"{c}20"] = tmpl.format(r=f"{c}19")
    check(label, grade_row20_budget_conversion_v2(ws)[1], want)

print("MA3 · Analysis G36:G37")
check("=I18-1*(I20) lower", _bound_credit("=I18-1*(I20)", True), 1.0)
check("=I18+1*(I20) upper", _bound_credit("=I18+1*(I20)", False), 1.0)
wb = Workbook()
ws = wb.active
ws["G36"] = "=AVERAGE(D14:D63)-_xlfn.STDEV.S(D14:D63)"   # as Excel saves it
ws["G37"] = "=AVERAGE(D14:D63)+_xlfn.STDEV.S(D14:D63)"
strip_future_prefixes(wb)
check("=AVERAGE(D14:D63)-STDEV.S(D14:D63) (saved with _xlfn.)", check_empirical_rule(ws)[0], 6.0)

print("MA3 · Written analysis mentioning 68%")
ws = Workbook().active
ws["C41"] = ("According to the empirical rule, approximately 68% of the students should expect "
             "to see an increase in their scores between 0.91 and 4.84 points.")
check("answer containing '68%' is found", check_written_analysis(ws)[1][0][0], "WRITTEN_FOUND")

print("MA3 · Visualization")
check("G: =COUNTIFS(B12:B61,D28:D38,...) rejected",
      _check_frequency_formula("=COUNTIFS(B12:B61, D28:D38, B12:B61, E28:E38)", spill=True), False)
check("G: comparison COUNTIFS accepted",
      _check_frequency_formula('=COUNTIFS($B$12:$B$61,">="&D30,$B$12:$B$61,"<"&E30)', 30), True)
check("H: dragged =G29/SUM(G29:G39) rejected", _check_relative_freq_formula("=G29/SUM(G29:G39)", 29), False)
check("H: =G29/SUM(G$28:G$38) accepted", _check_relative_freq_formula("=G29/SUM(G$28:G$38)", 29), True)
check("H: dragged =G29/COUNT(B14:B62) rejected", _check_relative_freq_formula("=G29/COUNT(B14:B62)", 29), False)
check("H: =G29/COUNT(B$12:B$61) accepted", _check_relative_freq_formula("=G29/COUNT(B$12:B$61)", 29), True)
check("E22: =MIN(Analysis!D14:D63)-0.1 accepted", _check_min_formula("=MIN(Analysis!D14:D63)-0.1"), True)

ws = Workbook().active
for r in range(28, 39):
    ws[f"F{r}"] = f"=(D{r}+E{r})/2"
    ws[f"G{r}"] = f'=COUNTIFS($B$12:$B$61,">="&D{r},$B$12:$B$61,"<"&E{r})'
    ws[f"H{r}"] = f"=G{r}/SUM(G$28:G$38)"
check("fully correct table scores 18", check_freq_dist_values(ws)[0], 18.0)

print()
if failures:
    print(f"{len(failures)} FAILED: {failures}")
    sys.exit(1)
print("All reported issues stay fixed.")
